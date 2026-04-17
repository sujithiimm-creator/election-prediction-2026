"""
pipeline/output_generator.py
------------------------------
Generates the final submission Excel file in competition format.

Also generates:
  - Methodology note (required by competition rules)
  - Confidence report (internal use)
  - State-wise summary
"""

import os
import pandas as pd
import numpy as np
from datetime import datetime
from typing import Dict, List, Optional
from src.utils.helpers import get_logger, ensure_dir

logger = get_logger(__name__)


class SubmissionGenerator:
    """
    Converts internal prediction DataFrame into competition submission format.

    Competition requirement (from T&C):
      - Excel file with predicted winning candidate/party per constituency
      - Methodology note PDF/Word (150+ words)
    """

    def __init__(self, output_dir: str = "outputs/predictions/"):
        self.output_dir = output_dir
        ensure_dir(output_dir)

    # ── Main: Generate submission file ────────────────────────────────────

    def generate_submission(
        self,
        predictions: pd.DataFrame,
        filename: str = "final_predictions.xlsx",
        include_confidence: bool = True,
    ) -> str:
        """
        Create the final submission Excel.
        One row per constituency, showing predicted winner + supporting data.
        """
        # Filter to only winner rows
        winners = predictions[predictions["predicted_winner"] == 1].copy()

        # Build submission DataFrame
        sub_cols = [
            "state", "constituency_name", "party",
            "predicted_vs", "pred_margin_after_rules",
        ]
        if include_confidence:
            sub_cols += ["confidence_level", "flag_for_review", "rule_applied"]

        available_cols = [c for c in sub_cols if c in winners.columns]
        submission = winners[available_cols].copy()

        # Clean up for submission
        submission = submission.rename(columns={
            "party":                   "predicted_winner_party",
            "predicted_vs":            "predicted_vote_share_pct",
            "pred_margin_after_rules": "predicted_margin_pct",
            "rule_applied":            "model_rule_applied",
        })

        submission["predicted_vote_share_pct"] = submission["predicted_vote_share_pct"].round(2)
        if "predicted_margin_pct" in submission.columns:
            submission["predicted_margin_pct"] = submission["predicted_margin_pct"].round(2)

        submission = submission.sort_values(["state", "constituency_name"])

        filepath = os.path.join(self.output_dir, filename)

        with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
            # Sheet 1: All states combined
            submission.to_excel(writer, sheet_name="All_States", index=False)

            # Sheet per state
            for state in submission["state"].unique():
                state_df = submission[submission["state"] == state]
                sheet_name = state.replace(" ", "_")[:31]   # Excel limit
                state_df.to_excel(writer, sheet_name=sheet_name, index=False)

            # Summary sheet
            summary = self._build_summary(submission, predictions)
            summary.to_excel(writer, sheet_name="Summary", index=False)

            # Format sheets
            self._format_excel(writer)

        logger.info(
            "Submission file saved: %s (%d constituencies across %d states)",
            filepath,
            len(submission),
            submission["state"].nunique(),
        )
        return filepath

    # ── Summary builder ───────────────────────────────────────────────────

    def _build_summary(
        self, submission: pd.DataFrame, full_predictions: pd.DataFrame
    ) -> pd.DataFrame:
        records = []
        for state in submission["state"].unique():
            state_sub = submission[submission["state"] == state]
            n_total = len(state_sub)
            n_flagged = state_sub.get("flag_for_review", pd.Series(False, index=state_sub.index)).sum()
            n_high_conf = (state_sub.get("confidence_level", "medium") == "high").sum()

            # Party seat count
            party_seats = (
                state_sub.groupby("predicted_winner_party")
                .size()
                .sort_values(ascending=False)
                .head(3)
            )
            top_parties = ", ".join(
                [f"{p}: {n}" for p, n in party_seats.items()]
            )

            records.append({
                "state": state,
                "total_constituencies": n_total,
                "high_confidence": int(n_high_conf),
                "needs_review": int(n_flagged),
                "top_parties_seats": top_parties,
            })

        return pd.DataFrame(records)

    # ── Excel formatting ──────────────────────────────────────────────────

    @staticmethod
    def _format_excel(writer):
        """Apply basic formatting to make the Excel readable."""
        from openpyxl.styles import PatternFill, Font, Alignment
        
        red_fill    = PatternFill("solid", fgColor="FFE0E0")
        green_fill  = PatternFill("solid", fgColor="E0FFE0")
        yellow_fill = PatternFill("solid", fgColor="FFFACD")
        header_font = Font(bold=True)

        for sheet_name, ws in writer.sheets.items():
            # Bold headers
            for cell in ws[1]:
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

            # Auto-width columns
            for col in ws.columns:
                max_len = max(
                    (len(str(cell.value)) for cell in col if cell.value),
                    default=10
                )
                ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

            # Highlight flagged rows (flag_for_review = True)
            flag_col_idx = None
            for i, cell in enumerate(ws[1], 1):
                if cell.value == "flag_for_review":
                    flag_col_idx = i
                    break

            if flag_col_idx:
                for row in ws.iter_rows(min_row=2):
                    flag_val = row[flag_col_idx - 1].value
                    if flag_val is True or flag_val == "True":
                        for cell in row:
                            cell.fill = yellow_fill


# ── Methodology note generator ────────────────────────────────────────────────

def generate_methodology_note(
    metrics: Optional[Dict] = None,
    output_path: str = "outputs/reports/methodology_note.txt",
) -> str:
    """
    Auto-generate the required methodology note (150+ words).
    Customise this with your actual approach details.
    """
    ensure_dir(os.path.dirname(output_path))

    acc_str = ""
    if metrics and "overall_accuracy" in metrics:
        acc_str = f"Backtesting on the 2021 election produced an overall winner accuracy of {metrics['overall_accuracy']*100:.1f}%."

    note = f"""
METHODOLOGY NOTE — India Predicts 2026
Election Prediction Challenge | Data Science Academy
Submission Date: {datetime.now().strftime('%B %d, %Y')}

─────────────────────────────────────────────────────
APPROACH OVERVIEW
─────────────────────────────────────────────────────

This submission uses a hybrid prediction system combining historical 
data modelling, rule-based logic, and expert domain knowledge.

DATA SOURCES
Historical election results from the Election Commission of India (ECI) 
for the years 2011, 2016, and 2021 were used across all five states. 
Candidate-level vote counts, party affiliations, constituency totals, 
and elector data were extracted and standardised from ECI statistical 
reports. Additional alliance information was sourced from publicly 
available news coverage and MyNeta.info candidate profiles.

FEATURE ENGINEERING
For each constituency × party pair, the following features were computed 
using only data available before the target election (strict no-leakage 
protocol): historical vote share at t-1, t-2, t-3; vote share trend 
(linear slope); incumbency flags (party and candidate); margin of victory 
categories (safe/swing/tossup); constituency volatility; stronghold 
indicators (party wins across 3 elections); turnout change; and 
alliance-based vote transfer estimates.

Alliance relationships were inferred dynamically from co-occurrence 
patterns: if two parties consistently avoided contesting the same 
constituency, they were classified as allies with an estimated vote 
transfer coefficient.

MACHINE LEARNING MODEL
A LightGBM gradient boosting regressor was trained to predict vote 
share per party per constituency. The winner was then derived as the 
party with the highest predicted vote share. Training used a time-aware 
walk-forward cross-validation strategy: features computed from elections 
t-2 and t-1, validated on election t.

RULE-BASED ADJUSTMENTS
Three primary rule layers were applied on top of ML predictions:
(1) Stronghold protection: boost for parties winning 2+ of last 3 
elections in safe seats; (2) Anti-incumbency dampening: 10% reduction 
for volatile seats held by the state's ruling party; (3) Alliance 
vote transfer: computed transfer-in scores based on ally historical 
vote shares.

HUMAN EXPERT REVIEW
Low-confidence constituencies (predicted margin < 6%) were flagged for 
expert review. Tamil Nadu constituency predictions were manually reviewed 
using domain knowledge of local candidate strength, caste dynamics, 
and 2026 political developments.

{acc_str}

LIMITATIONS
Historical patterns may not capture sudden political realignments or 
candidate-specific events occurring close to polling. Alliance data 
is subject to last-minute changes. Rural booth-level granularity was 
not available for all states.

─────────────────────────────────────────────────────
Word count: ~350 words
─────────────────────────────────────────────────────
""".strip()

    with open(output_path, "w") as f:
        f.write(note)
    logger.info("Methodology note saved to %s", output_path)
    return output_path
